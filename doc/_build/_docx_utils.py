"""Shared helpers for building LLD .docx files (one per section batch).

Usage from a sibling script:
    from _docx_utils import LLDBuilder
    b = LLDBuilder("Anomaly Agent — LLD Sections 1–5")
    b.h1("1. Product Vision & Problem Statement")
    b.h2("1.1 Mission")
    b.p("...")
    b.status_table([
        ("Mission statement", "IMPLEMENTED", "doc/lld.md §1.1"),
    ])
    b.save("/path/to/output.docx")
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches


STATUS_COLORS = {
    "IMPLEMENTED": RGBColor(0x1F, 0x77, 0x44),
    "PARTIAL":     RGBColor(0xB7, 0x6E, 0x00),
    "PLANNED":     RGBColor(0x8C, 0x1A, 0x1A),
    "NOT USED":    RGBColor(0x55, 0x55, 0x55),
}


def _set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


class LLDBuilder:
    def __init__(self, title: str) -> None:
        self.doc = Document()
        # base style
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)

        # title page heading
        t = self.doc.add_heading(title, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.LEFT

        sub = self.doc.add_paragraph()
        sub.add_run(
            "Branch: claude/anomaly-agent-frontend-s9ygV   ·   "
            "Source template: doc/lld.md   ·   "
            "Status legend: IMPLEMENTED / PARTIAL / PLANNED"
        ).italic = True

        # legend
        legend = self.doc.add_paragraph()
        legend.add_run(
            "IMPLEMENTED = wired end-to-end on this branch and runs against real "
            "services (LLM + MySQL + SSH) when reachable. "
        )
        legend.add_run(
            "PARTIAL = present but limited (e.g. regex-only, hard-coded table, no "
            "persistence). "
        )
        legend.add_run(
            "PLANNED = not in code today; documented for direction."
        )

        self.doc.add_paragraph()  # spacer

    # ------------------------------------------------------------------ headings
    def h1(self, text: str) -> None:
        self.doc.add_heading(text, level=1)

    def h2(self, text: str) -> None:
        self.doc.add_heading(text, level=2)

    def h3(self, text: str) -> None:
        self.doc.add_heading(text, level=3)

    # ------------------------------------------------------------------ paragraphs
    def p(self, text: str) -> None:
        self.doc.add_paragraph(text)

    def bullets(self, items: list[str]) -> None:
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def numbered(self, items: list[str]) -> None:
        for it in items:
            self.doc.add_paragraph(it, style="List Number")

    def code(self, text: str) -> None:
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        # subtle gray background
        rPr = run._r.get_or_add_rPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F4F4F4")
        rPr.append(shd)

    # ------------------------------------------------------------------ tables
    def status_table(
        self,
        rows: list[tuple[str, str, str]],
        headers: tuple[str, str, str] = ("Topic", "Status", "Where it lives / Notes"),
    ) -> None:
        """3-column table: topic, status (color-coded), notes/source."""
        table = self.doc.add_table(rows=1 + len(rows), cols=3)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT

        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            run.bold = True
            _set_cell_shading(cell, "DDDDDD")

        for r, (topic, status, notes) in enumerate(rows, start=1):
            row = table.rows[r].cells
            row[0].text = topic
            status_cell = row[1]
            status_cell.text = ""
            run = status_cell.paragraphs[0].add_run(status)
            run.bold = True
            color = STATUS_COLORS.get(status, RGBColor(0, 0, 0))
            run.font.color.rgb = color
            row[2].text = notes

        # Column widths
        widths = (Inches(2.0), Inches(1.1), Inches(3.6))
        for col, w in zip(table.columns, widths):
            for cell in col.cells:
                cell.width = w

        self.doc.add_paragraph()  # spacer

    def kv_table(self, rows: list[tuple[str, str]]) -> None:
        table = self.doc.add_table(rows=len(rows), cols=2)
        table.style = "Light Grid Accent 1"
        for r, (k, v) in enumerate(rows):
            table.rows[r].cells[0].text = k
            table.rows[r].cells[1].text = v
            for run in table.rows[r].cells[0].paragraphs[0].runs:
                run.bold = True
        widths = (Inches(2.0), Inches(4.7))
        for col, w in zip(table.columns, widths):
            for cell in col.cells:
                cell.width = w
        self.doc.add_paragraph()

    # ------------------------------------------------------------------ save
    def save(self, path: str | Path) -> None:
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
